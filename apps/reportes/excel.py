import datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY_FORMAT = '"$" #,##0.00'


def excel_response(filename, sheet_title, headers, rows, money_cols=None):
    """Genera un .xlsx descargable. Las columnas Decimal se formatean como pesos."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for row in rows:
        ws.append([v if v is None or isinstance(v, (int, float, Decimal, datetime.date, str)) else str(v) for v in row])
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 50)
        for c in ws[letter][1:]:
            if isinstance(c.value, Decimal) or (money_cols and col_idx - 1 in money_cols):
                c.number_format = MONEY_FORMAT
            elif isinstance(c.value, datetime.date):
                c.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}_{datetime.date.today():%Y%m%d}.xlsx"'
    return resp
